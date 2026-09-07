"""Every supported upload format, including the ones that need a real file."""

from __future__ import annotations

import io

import pytest

from app.core.errors import ValidationError
from app.knowledge.extraction import extension_of, extract
from app.knowledge.ingestion import clean


def _make_pdf(pages: list[str]) -> bytes:
    """Build a minimal PDF with a real text layer, by hand.

    Written as raw bytes rather than through a library's internals so it does
    not break when that library reorganises its private API.
    """
    objects: list[bytes] = []

    def add(body: bytes) -> int:
        objects.append(body)
        return len(objects)

    font_id = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    page_ids: list[int] = []
    # object 1 is the font, then two objects per page, then /Pages.
    pages_id_placeholder = len(pages) * 2 + 2

    for text in pages:
        escaped = text.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
        content = f"BT /F1 12 Tf 72 700 Td ({escaped}) Tj ET".encode("latin-1")
        stream_id = add(b"<< /Length %d >>\nstream\n%s\nendstream" % (len(content), content))
        page_id = add(
            b"<< /Type /Page /Parent %d 0 R /MediaBox [0 0 612 792] "
            b"/Contents %d 0 R /Resources << /Font << /F1 %d 0 R >> >> >>"
            % (pages_id_placeholder, stream_id, font_id)
        )
        page_ids.append(page_id)

    kids = b" ".join(b"%d 0 R" % pid for pid in page_ids)
    pages_id = add(b"<< /Type /Pages /Kids [%s] /Count %d >>" % (kids, len(page_ids)))
    assert pages_id == pages_id_placeholder, "object numbering drifted"
    catalog_id = add(b"<< /Type /Catalog /Pages %d 0 R >>" % pages_id)

    out = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % index + body + b"\nendobj\n"
    xref_at = len(out)
    out += b"xref\n0 %d\n" % (len(objects) + 1)
    out += b"0000000000 65535 f \n"
    for offset in offsets[1:]:
        out += b"%010d 00000 n \n" % offset
    out += b"trailer\n<< /Size %d /Root %d 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objects) + 1,
        catalog_id,
        xref_at,
    )
    return bytes(out)


def _make_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    import docx

    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        word_table = document.add_table(rows=len(table), cols=len(table[0]))
        for row_index, row in enumerate(table):
            for column_index, value in enumerate(row):
                word_table.cell(row_index, column_index).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class TestExtensions:
    @pytest.mark.parametrize(
        "filename,expected",
        [("a.PDF", ".pdf"), ("b.tar.gz", ".gz"), ("c.markdown", ".markdown"), ("noext", "")],
    )
    def test_extension_detection(self, filename, expected):
        assert extension_of(filename) == expected


class TestPdf:
    def test_a_pdf_with_a_text_layer_is_extracted_page_by_page(self):
        data = _make_pdf(["First page about VAT.", "Second page about ZUS."])
        result = extract(data, "report.pdf")
        assert "VAT" in result.text and "ZUS" in result.text
        assert "[page 1]" in result.text and "[page 2]" in result.text
        assert result.page_count == 2

    def test_a_scanned_pdf_with_no_text_layer_is_an_explicit_error(self):
        """A silently empty document is worse than a refused upload."""
        from pypdf import PdfWriter

        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)
        buffer = io.BytesIO()
        writer.write(buffer)
        with pytest.raises(ValidationError, match="OCR"):
            extract(buffer.getvalue(), "scan.pdf")

    def test_a_corrupt_pdf_is_an_error_not_a_crash(self):
        with pytest.raises(ValidationError, match="could not read"):
            extract(b"%PDF-1.4 this is not really a pdf", "broken.pdf")


class TestDocx:
    def test_paragraphs_and_tables_are_extracted(self):
        data = _make_docx(
            ["The VAT rate is 23%.", "Reduced rates apply to listed goods."],
            table=[["Rate", "Applies to"], ["23%", "standard"], ["8%", "construction"]],
        )
        result = extract(data, "rates.docx")
        assert "The VAT rate is 23%." in result.text
        assert "Rate | Applies to" in result.text
        assert result.meta["tables"] == 1

    def test_an_empty_docx_is_refused(self):
        with pytest.raises(ValidationError, match="no readable text"):
            extract(_make_docx([]), "empty.docx")

    def test_a_corrupt_docx_is_an_error_not_a_crash(self):
        with pytest.raises(ValidationError, match="could not read"):
            extract(b"PK\x03\x04 not really a docx", "broken.docx")


class TestCsvAndJson:
    def test_a_semicolon_delimited_csv_is_detected(self):
        result = extract(b"name;amount\nZUS;5203.80\nPIT;1200.00", "x.csv")
        assert "name | amount" in result.text and result.meta["columns"] == 2

    def test_an_empty_csv_is_refused(self):
        with pytest.raises(ValidationError):
            extract(b"", "x.csv")

    def test_jsonl_is_extracted_record_by_record(self):
        result = extract(b'{"a":1}\n{"a":2}\n', "x.jsonl")
        assert result.meta["records"] == 2

    def test_invalid_jsonl_names_the_line(self):
        with pytest.raises(ValidationError, match="line 2"):
            extract(b'{"a":1}\n{bad}\n', "x.jsonl")


class TestEncoding:
    def test_utf8_with_a_bom_is_handled(self):
        assert "Zażółć" in extract("Zażółć gęślą jaźń".encode("utf-8-sig"), "pl.txt").text

    def test_a_windows_1250_file_is_handled(self):
        assert extract("Zażółć".encode("cp1250"), "pl.txt").text

    def test_undecodable_bytes_do_not_crash(self):
        assert extract(b"\xff\xfe\x00\x01 some text", "x.txt").text


class TestCleaning:
    def test_control_characters_and_runs_of_whitespace_are_removed(self):
        # The null byte is removed, which joins "a" and "b" — that is the
        # intended behaviour: a control character is noise, not a separator.
        assert clean("a\x00b   c\n\n\n\nd\r\ne") == "ab c\n\nd\ne"

    def test_cleaning_an_empty_string_is_safe(self):
        assert clean("") == "" and clean(None) == ""
